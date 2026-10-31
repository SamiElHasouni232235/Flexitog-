"""RFQ Quick Quote MENA-TR v2.3 (assets/rfq_v23.js): parsing, validation, costs per option, scoring, red flags and
follow-up questions, run in node on real workbooks.

Test 1 and Test 4 need the provider files in tests/fixtures/ under their own names (they hold a provider's real
quote, so they are not in the repository); without them those tests skip. Every other test runs on the bundled
blank template, filled in by openpyxl here, and on the v1 dummy fixture.
"""
import datetime
import json
import shutil
import subprocess
from pathlib import Path

import openpyxl
import pytest

from flexitog.dashboard import ASSETS, RFQ_LEGACY_TEMPLATE, RFQ_TEMPLATE, rfq_sources

FIX = Path(__file__).parent / "fixtures"
ARAMEX_V23 = FIX / "FlexiTog RFQ Quick Quote MENA-TR v2.3 - Aramex test fill.xlsx"
ARAMEX_V1 = FIX / "Transport Volume Scenario RFQ FlexiTog 20271029 - Aramex reply.xlsx"
DUMMY_V1 = FIX / "Sample_RFQ_Response_DUMMY.xlsx"
LIBS = [str(ASSETS / "vendor" / "xlsx.full.min.js"), str(ASSETS / "rfq_core.js"), str(ASSETS / "rfq_v23.js")]

pytestmark = pytest.mark.skipif(shutil.which("node") is None, reason="node is not installed")


def run(body: str, tmp_path, *files):
    """Run JS with X (SheetJS), R (rfq core), V (rfq v2.3), fs, F[] (file paths) and SRC (template key -> cell maps)."""
    src = tmp_path / "sources.json"
    src.write_text(json.dumps({"v23": rfq_sources(RFQ_TEMPLATE), "v1": rfq_sources(RFQ_LEGACY_TEMPLATE)}))
    script = tmp_path / "t.js"
    script.write_text(f"""const X = require({json.dumps(LIBS[0])}), R = require({json.dumps(LIBS[1])}), V = require({json.dumps(LIBS[2])}), fs = require('fs');
const F = {json.dumps([str(f) for f in files])}, SRC = JSON.parse(fs.readFileSync({json.dumps(str(src))}, 'utf8'));
const load = (f, kind) => R.parseWorkbook(X.read(fs.readFileSync(f), {{cellDates: true}}), X, {{sources: SRC[kind || 'v23']}});
const evalFile = (f, settings) => {{ const P = load(f), values = V.valuesOf(P), E = V.evaluate(values, settings || null);
  const S = V.scoreAll([{{id: 'x', Q: E.Q, A: E.A}}], settings || null)[0]; return Object.assign(E, {{P, values, S, follow: V.followUps(E.Q, E.A, S, E.flags, settings || null)}}); }};
const optOut = O => ({{state: O.state, cost: O.cost, covered: O.covered, shipments: O.shipments, onTime: O.onTime, extrapolated: O.extrapolated, tag: O.tag}});
const summary = E => ({{version: E.Q.version, inputCheck: E.Q.inputCheck, invalid: E.Q.invalid, shipments: E.Q.shipments.length,
  opt: Object.fromEntries(['opt1', 'opt2', 'opt3'].map(o => [o, {{A: optOut(E.A[o].A), B: optOut(E.A[o].B)}}])),
  ship: Object.fromEntries(E.Q.shipments.map(s => [s.id, {{lane: s.lane, band: s.band, opt1: s.opt1}}])),
  areas: Object.fromEntries(E.S.areas.map(a => [a.key, a.score])), total: E.S.total, option: E.S.option, missing: E.S.missing,
  flags: E.flags.map(f => f.id), follow: E.follow}});
let out = null;
(async () => {{ {body}
  process.stdout.write(JSON.stringify(out)); }})().catch(e => {{ console.error(e); process.exit(1); }});
""")
    r = subprocess.run(["node", str(script)], capture_output=True, text=True, timeout=120)
    assert r.returncode == 0, r.stderr
    return json.loads(r.stdout)


def fill(path, *, incl_cell=None, opt3=True, warehouse=False):
    """A filled copy of the blank template, made like a provider would: prices on the Test Shipments sheet, answers on
    Questions and Warehouse. openpyxl keeps the Export formulas without cached results, so the parser reads the input
    cells, as it does for any file not saved in Excel."""
    wb = openpyxl.load_workbook(RFQ_TEMPLATE)
    ts, qs, wh = wb["Test Shipments"], wb["Questions"], wb["Warehouse"]
    price = {"S": 60, "M": 240, "L": 900, "XL": 3100, "XXL": 7500}
    for r in range(19, 52):
        tid, lane, size = ts.cell(r, 1).value, ts.cell(r, 2).value, ts.cell(r, 5).value
        if tid in ("T16", "T33"):
            continue                                            # not offered: L05 S and L09 L
        ts.cell(r, 9).value = "Express" if size in ("S", "M") else "Air freight"
        ts.cell(r, 10).value = price[size] * (1.2 if lane in ("L07", "L08") else 1)
        ts.cell(r, 11).value = 2 if lane == "L02" and size == "S" else 4
        if opt3 and lane in ("L05", "L06"):
            ts.cell(r, 14).value, ts.cell(r, 15).value, ts.cell(r, 16).value = "Road groupage", price[size] * 0.7, 6
        if warehouse and lane == "L02":
            ts.cell(r, 12).value, ts.cell(r, 13).value = 20, 1
    ts.cell(19, 17).value = "Volumetric 1:5000 applies"
    if incl_cell:
        ts[incl_cell] = "incl"
    answers = {6: "Example Co", 7: "A. Person", 8: "a@example.com", 9: "+31 6 1234", 10: datetime.datetime(2026, 12, 31), 11: 0.3,
               12: "Some (list in Comment)", 13: 30, 15: datetime.time(15, 30), 16: "Yes from Helmond", 17: 2, 18: "Partly guaranteed",
               20: 0.96, 21: 0, 22: "Both", 23: "One of the two", 24: "Yes", 25: "Yes", 26: "Only ISO 9001", 28: "All 9 countries",
               29: "All 9 countries", 30: "Some (list in Comment)", 31: "No", 32: "All 9 countries", 33: "No", 34: "Yes against a fee", 35: 0.002, 36: "Yes"}
    for row, v in answers.items():
        qs.cell(row, 3).value = v
    qs.cell(11, 4).value = "Fuel 30% of November 2026"
    qs.cell(12, 4).value = "Per piece surcharge above 70 kg"
    wh.cell(6, 3).value = "Dubai, AE"
    if warehouse:
        for row, v in {10: 12, 11: 8, 12: 0.5, 13: 3, 15: 250, 18: 180}.items():   # W04 W05 W06 W07 W09 W11
            wh.cell(row, 3).value = v
    wb.save(path)
    return path


# ------------------------------------------------------------------ Test 1: the Aramex test fill
@pytest.mark.skipif(not ARAMEX_V23.is_file(), reason=f"put {ARAMEX_V23.name} in tests/fixtures to run Test 1")
def test_1_aramex_v23(tmp_path):
    s = run("out = summary(evalFile(F[0]));", tmp_path, ARAMEX_V23)
    o1A, o1B = s["opt"]["opt1"]["A"], s["opt"]["opt1"]["B"]
    assert s["version"] == "RFQ-MENA-TR-QUICK v2.3"
    assert o1A["cost"] == pytest.approx(188208.28, abs=0.01) and (o1A["covered"], o1A["shipments"]) == (129, 141)
    assert o1A["extrapolated"] == 0 and o1A["onTime"] == 2
    assert o1B["cost"] == pytest.approx(446261.23, abs=0.01) and (o1B["covered"], o1B["shipments"]) == (301, 329) and o1B["onTime"] == 4
    assert s["ship"]["T05"]["lane"] == "L01" and s["ship"]["T05"]["opt1"]["price"] == pytest.approx(17894.40, abs=0.005)
    assert s["ship"]["T16"]["lane"] == "L05" and s["ship"]["T16"]["opt1"]["status"] == "not_offered"
    assert s["opt"]["opt2"]["A"]["state"] == "incomplete" and s["opt"]["opt3"]["A"]["state"] == "not_quoted"
    a = s["areas"]
    assert round(a["cost"], 2) == 4.00 and round(a["service"], 2) == 1.60 and round(a["quality"], 2) == 3.41
    assert round(a["options"], 2) == 1.56 and round(s["total"], 2) == 2.64
    assert sorted(set(s["flags"])) == ["F1", "F2", "F3", "F4", "F5", "F8"]


# ------------------------------------------------------------------ Test 2: the blank template
def test_2_blank_template_offers_nothing(tmp_path):
    s = run("out = summary(evalFile(F[0]));", tmp_path, RFQ_TEMPLATE)
    assert s["version"] == "RFQ-MENA-TR-QUICK v2.3" and s["shipments"] == 33 and s["invalid"] == []
    assert s["inputCheck"] == "No problems found"
    for o in ("opt1", "opt2", "opt3"):
        for sc, n in (("A", 141), ("B", 329)):
            x = s["opt"][o][sc]
            assert x["state"] == "not_quoted" and x["cost"] is None and x["covered"] == 0 and x["shipments"] == n
    assert all(v["opt1"]["status"] == "not_offered" for v in s["ship"].values())
    assert s["total"] == 0 and "F6" in s["flags"]


# ------------------------------------------------------------------ Test 3: text in a price cell
def test_3_text_in_a_price_is_invalid_never_zero(tmp_path):
    clean, dirty = fill(tmp_path / "clean.xlsx"), fill(tmp_path / "incl.xlsx", incl_cell="J20")   # J20 = T02 Option 1 price
    s = run("out = [summary(evalFile(F[0])), summary(evalFile(F[1]))];", tmp_path, clean, dirty)
    c, d = s
    assert c["invalid"] == [] and d["invalid"] == [{"key": "ship.T02.opt1_price_eur", "value": "incl", "reason": "text"}]
    assert d["ship"]["T02"]["opt1"]["price"] is None and d["ship"]["T02"]["opt1"]["status"] == "not_offered"
    # T02 is L01 B2: 8 Scenario A shipments drop out of coverage and cost, nothing is counted at 0
    assert c["opt"]["opt1"]["A"]["covered"] - d["opt"]["opt1"]["A"]["covered"] == 8
    assert c["opt"]["opt1"]["A"]["cost"] - d["opt"]["opt1"]["A"]["cost"] == pytest.approx(8 * 240)


@pytest.mark.skipif(not ARAMEX_V23.is_file(), reason=f"put {ARAMEX_V23.name} in tests/fixtures to run Test 3 on the Aramex file")
def test_3_aramex_copy_with_incl(tmp_path):
    wb = openpyxl.load_workbook(ARAMEX_V23)
    ts = wb["Test Shipments"]
    row = next(r for r in range(19, 52) if ts.cell(r, 1).value == "T05")
    ts.cell(row, 10).value = "incl"
    wb.save(tmp_path / "aramex_incl.xlsx")
    s = run("out = summary(evalFile(F[0]));", tmp_path, tmp_path / "aramex_incl.xlsx")
    assert {"key": "ship.T05.opt1_price_eur", "value": "incl", "reason": "text"} in s["invalid"]
    assert s["ship"]["T05"]["opt1"]["price"] is None
    assert s["opt"]["opt1"]["A"]["covered"] == 129 - 2           # L01 B5 has 2 Scenario A shipments
    assert s["opt"]["opt1"]["A"]["cost"] == pytest.approx(188208.28 - 2 * 17894.40, abs=0.01)


# ------------------------------------------------------------------ Test 4: a v1 answer stays legacy
def _legacy(tmp_path, f):
    return run("""const P0 = R.parseWorkbook(X.read(fs.readFileSync(F[0]), {cellDates: true}), X, {});
      const P = load(F[0], 'v1'), check = R.validate(P, null);
      out = {kind: V.versionKind(P0.version), version: P.version, ok: check.ok, errors: check.errors, costB: R.summarize(R.effectiveValues(P)).costB};""", tmp_path, f)


def test_4_v1_dummy_imports_as_legacy(tmp_path):
    s = _legacy(tmp_path, DUMMY_V1)
    assert s["kind"] == "legacy" and s["version"] == "RFQ-MENA-TR v1.0" and s["ok"] and s["costB"] is not None


@pytest.mark.skipif(not ARAMEX_V1.is_file(), reason=f"put {ARAMEX_V1.name} in tests/fixtures to run Test 4")
def test_4_aramex_v1_imports_as_legacy(tmp_path):
    s = _legacy(tmp_path, ARAMEX_V1)
    assert s["kind"] == "legacy" and s["ok"], s["errors"]


def test_version_handling():
    out = subprocess.run(["node", "-e", f"""const V = require({json.dumps(LIBS[2])});
      console.log(JSON.stringify(['RFQ-MENA-TR-QUICK v2.3', 'RFQ-MENA-TR v1.0', 'RFQ-MENA-TR-QUICK v9.9', ''].map(V.versionKind)));"""],
                         capture_output=True, text=True, check=True).stdout
    assert json.loads(out) == ["v23", "legacy", "unknown", "unknown"]


# ------------------------------------------------------------------ costs, scores, flags on a filled copy
def test_costs_scores_and_flags_by_hand(tmp_path):
    s = run("out = summary(evalFile(F[0]));", tmp_path, fill(tmp_path / "f.xlsx"))
    o1 = s["opt"]["opt1"]["A"]
    # T16 (L05 B1, 6 shipments) not offered; every other A band has a test shipment and a price
    assert (o1["covered"], o1["shipments"], o1["extrapolated"]) == (135, 141, 0) and o1["tag"] == "derived"
    assert s["option"] == "opt1"                                  # Option 3 covers 2 lanes only: not complete
    a = s["areas"]
    assert a["cost"] == pytest.approx(4)
    # on time 43 of 141 (30%) -> 1; Q10 Yes from Helmond 4; Q09 15:30 -> 3; Q11 2 days -> 3; Q12 Partly 2
    assert a["service"] == pytest.approx((1 + 4 + 3 + 3 + 2) / 5)
    # Q13 96% -> 3 x 0.75; Q14 0% -> 4 x 0.75; Q15 Both 4; Q16 One 2; Q17 Yes 4 twice; Q18 Yes 4; Q19 Only ISO 9001 2
    assert a["quality"] == pytest.approx((2.25 + 3 + 4 + 2 + 8 + 4 + 2) / 8)
    # Q20 4, Q21 4, Q22 Some 2, Q23 No 0, Q24 4, Q25 No 0, Q26 against a fee 3, Option 2 only W01 1, Option 3 on 2 lanes 2
    assert a["options"] == pytest.approx(20 / 9)
    assert s["total"] == pytest.approx((4 + 2.6 + 3.15625 + 20 / 9) / 4)
    # valid until 2026-12-31 (F1), per piece / 70 kg (F2), 1:5000 (F3), fuel for November only (F4), no cached input check (F7), damage 0 (F8)
    assert sorted(set(s["flags"])) == ["F1", "F2", "F3", "F4", "F7", "F8"]
    assert any(q.startswith("Can you quote Option 1 for L05 Türkiye B1") for q in s["follow"])
    assert any(q.startswith("Option 2 is incomplete: please fill in") for q in s["follow"])


def test_flag_f5_and_f6(tmp_path):
    # every GCC transit above 2 days with Q10 Yes -> F5; L09 with no price at all -> F6
    out = run("""const v = V.valuesOf(load(F[0]));
      Object.keys(v).filter(k => /^ship\\.T(0[1-9]|1[0-5])\\.opt1_transit_days$/.test(k)).forEach(k => { v[k] = 3; });
      Object.keys(v).filter(k => /^ship\\.T3[1-3]\\.opt[13]_price_eur$/.test(k)).forEach(k => { v[k] = ''; });
      out = V.evaluate(v, null).flags.map(f => f.id + ':' + f.key);""", tmp_path, fill(tmp_path / "f.xlsx"))
    assert "F5:q.Q10.answer" in out and "F6:ship.L09" in out


def test_option2_cost_per_order_by_hand(tmp_path):
    out = run("""const v = {'meta.rfq_version': V.VERSION, 'meta.input_check': V.OK_CHECK};
      [['T01', 'B1', 8], ['T02', 'B2', 52], ['T03', 'B3', 261]].forEach(([t, b, kg]) => Object.assign(v, {[`ship.${t}.lane`]: 'L01', [`ship.${t}.band`]: b,
        [`ship.${t}.gross_kg`]: kg, [`ship.${t}.opt1_price_eur`]: kg * 2, [`ship.${t}.opt1_transit_days`]: 3, [`ship.${t}.opt2_delivery_price_eur`]: 10, [`ship.${t}.opt2_transit_days`]: 1}));
      Object.assign(v, {'wh.W01.answer': 'Dubai', 'wh.W04.answer': 12, 'wh.W05.answer': 8, 'wh.W06.answer': 0.5, 'wh.W07.answer': 3, 'wh.W09.answer': 250, 'wh.W11.answer': 180});
      const zero = {B1: 0, B2: 0, B3: 0, B4: 0, B5: 0}, vol = Object.fromEntries(Object.keys(V.LANES).map(L => [L, L === 'L01' ? {B1: 10, B2: 5, B3: 2, B4: 1, B5: 0} : zero]));
      const Q = V.quote(v), o2 = V.optionCost(Q, 'opt2', 'A', {avgPallets: 20, volumes: {A: vol}}), o1 = V.optionCost(Q, 'opt1', 'A', {avgPallets: 20, volumes: {A: vol}});
      const none = V.optionCost(Q, 'opt2', 'A', {volumes: {A: vol}});
      out = {state: o2.state, fixed: o2.fixedShare, b1: o2.cells.find(c => c.band === 'B1').price, b4: o2.cells.find(c => c.band === 'B4'),
        o1b4: o1.cells.find(c => c.band === 'B4'), none: [none.state, none.inputs.needs], noPrices: V.opt2Inputs(V.quote(Object.assign({}, v, {'ship.T01.opt2_delivery_price_eur': '', 'ship.T02.opt2_delivery_price_eur': '', 'ship.T03.opt2_delivery_price_eur': '', 'wh.W05.answer': ''})), {}).missing};""", tmp_path)
    fixed = (250 * 12 + 12 * 20 * 12) / 18                       # W09 x 12 + W04 x pallets x 12, over 18 orders a year
    assert out["state"] == "ok" and out["fixed"] == pytest.approx(fixed)
    assert out["b1"] == pytest.approx(10 + 3 + 0.5 * 6 + 0 + (180 + 8) / 250 * 8 + fixed)
    # B4 has no test shipment: the B3 delivery price per kg x 941 kg, tagged extrapolated
    assert out["b4"]["tag"] == "extrapolated" and out["b4"]["delivery"] == pytest.approx(10 / 261 * 941)
    assert out["o1b4"]["price"] == pytest.approx(522 / 261 * 941) and out["o1b4"]["tag"] == "extrapolated"
    assert out["none"] == ["incomplete", ["average pallets in stock"]]
    assert out["noPrices"] == ["ship.Txx.opt2_delivery_price_eur", "wh.W05"]


def test_values_conversion(tmp_path):
    out = run("""out = [V.convert(0.6666667, 'time', 'q.Q09.answer'), V.convert('16:00', 'time', 'q.Q09.answer'), V.convert(46387, 'date', 'q.Q05.answer'),
      V.convert(96, 'percent', 'q.Q13.answer'), V.convert(0.95, 'percent', 'q.Q13.answer'), V.convert('95%', 'percent', 'q.Q13.answer'),
      V.convert('incl', 'number', 'ship.T01.opt1_price_eur'), V.convert('1.234,50', 'number', 'ship.T01.opt1_price_eur'), V.convert('€ 80', 'number', 'ship.T01.opt1_price_eur')];""", tmp_path)
    assert out == ["16:00", "16:00", "2026-12-31", 0.96, 0.95, 0.95, "incl", 1234.5, 80]


def test_option2_complete_in_a_filled_file(tmp_path):
    out = run("""const S = {avgPallets: 20}, E = evalFile(F[0], S);
      out = {o2: optOut(E.A.opt2.A), item: E.S.areas.find(a => a.key === 'options').items.find(i => i.id === 'opt2').points};""",
              tmp_path, fill(tmp_path / "w.xlsx", warehouse=True))
    assert out["o2"]["state"] == "ok" and out["o2"]["covered"] == 36 and out["o2"]["tag"] == "assumption" and out["item"] == 4

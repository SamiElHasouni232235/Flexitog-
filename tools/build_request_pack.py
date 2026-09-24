"""Build the data request workbook: one sheet per dataset, headers matching the importer.

Run:  python tools/build_request_pack.py [output.xlsx]
"""
from __future__ import annotations

import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from flexitog import parameters as P  # noqa: E402
from flexitog.schema import ENTITIES  # noqa: E402
from flexitog.datarequest import DATASETS  # noqa: E402
from flexitog.seed import SEED  # noqa: E402

FONT = "Arial"
HEAD_FILL = PatternFill("solid", fgColor="1F3A55")
REQ_FILL = PatternFill("solid", fgColor="8C2F1B")
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")
EXAMPLE_FONT = Font(name=FONT, italic=True, color="7F7F7F", size=10)
THIN = Border(bottom=Side(style="thin", color="BFBFBF"))



def style_header(cell, required: bool = False):
    cell.font = Font(name=FONT, bold=True, color="FFFFFF", size=10)
    cell.fill = REQ_FILL if required else HEAD_FILL
    cell.alignment = Alignment(vertical="center", wrap_text=True)


def entity_sheet(wb, title: str, key: str):
    entity = ENTITIES[key]
    ws = wb.create_sheet(title)
    for col, f in enumerate(entity.fields, start=1):
        c = ws.cell(row=1, column=col, value=f.name)
        style_header(c, f.required)
        note = (f.description or f.name) + (f"\nAllowed: {', '.join(f.allowed)}" if f.allowed else "")
        note += f"\nType: {f.dtype}" + ("\nRequired" if f.required else "")
        c.comment = Comment(note, "FlexiTog simulator", width=260, height=110)
        ws.column_dimensions[get_column_letter(col)].width = max(14, len(f.name) + 4)
    example = (SEED.get(key) or [{}])[0]
    for col, f in enumerate(entity.fields, start=1):
        v = example.get(f.name)
        c = ws.cell(row=2, column=col, value=v)
        c.font = EXAMPLE_FONT
    ws.cell(row=2, column=len(entity.fields) + 1, value="EXAMPLE ROW: delete before upload").font = Font(
        name=FONT, bold=True, color="8C2F1B", size=10)
    for r in range(3, 203):
        for col in range(1, len(entity.fields) + 1):
            ws.cell(row=r, column=col).fill = INPUT_FILL
            ws.cell(row=r, column=col).font = Font(name=FONT, size=10)
    for col, f in enumerate(entity.fields, start=1):
        if f.allowed:
            dv = DataValidation(type="list", formula1='"' + ",".join(f.allowed) + '"', allow_blank=True)
            ws.add_data_validation(dv)
            dv.add(f"{get_column_letter(col)}3:{get_column_letter(col)}202")
    ws.freeze_panes = "A2"
    return ws, len(entity.fields)


def param_sheet(wb, title: str, name: str):
    df = P.default_table(name)
    ws = wb.create_sheet(title)
    cols = list(df.columns)
    for col, name_ in enumerate(cols, start=1):
        c = ws.cell(row=1, column=col, value=name_)
        style_header(c)
        ws.column_dimensions[get_column_letter(col)].width = 44 if name_ == "source_note" else max(13, len(name_) + 3)
    for r, row in enumerate(df.itertuples(index=False), start=2):
        for col, v in enumerate(row, start=1):
            c = ws.cell(row=r, column=col, value=None if (isinstance(v, float) and v != v) else v)
            c.font = Font(name=FONT, size=10, color="0000FF" if cols[col - 1] not in ("source", "source_note")
                          else "000000")
            c.fill = INPUT_FILL
            c.border = THIN
    src_col = get_column_letter(cols.index("source") + 1)
    dv = DataValidation(type="list", formula1='"placeholder,real"', allow_blank=False)
    ws.add_data_validation(dv)
    dv.add(f"{src_col}2:{src_col}{len(df) + 60}")
    ws.freeze_panes = "A2"
    return ws, len(df), src_col


def build(path: Path):
    wb = Workbook()
    ov = wb.active
    ov.title = "Overview"
    ov["A1"] = "FlexiTog route simulator: data request pack"
    ov["A1"].font = Font(name=FONT, bold=True, size=14, color="1F3A55")
    ov["A2"] = ("The simulator runs on dummy data today. Each sheet below collects one dataset that replaces "
                "placeholders with company data. Work top-down: High priority moves the results most.")
    ov["A2"].font = Font(name=FONT, size=10)
    legend = [
        ("How to fill", ""),
        ("Dark blue header", "column name the tool reads. Keep the header text unchanged."),
        ("Dark red header", "required column."),
        ("Yellow cells", "fill in. On parameter sheets, overwrite the placeholder value, set source to real "
                         "and write where the number came from in source_note."),
        ("Grey italic row 2", "example on master data sheets. Delete it before upload."),
        ("Upload", "save the sheet as its own file or upload the workbook and pick the sheet on the page "
                   "named in column H."),
    ]
    for i, (k, v) in enumerate(legend, start=4):
        ov.cell(row=i, column=1, value=k).font = Font(name=FONT, bold=True, size=10)
        ov.cell(row=i, column=2, value=v).font = Font(name=FONT, size=10)
    header_row = 11
    heads = ["#", "Dataset (sheet)", "Priority", "Why it matters", "Where to get it", "Owner", "Due date",
             "Upload in tool", "Status", "Rows filled", "Rows marked real"]
    for col, h in enumerate(heads, start=1):
        style_header(ov.cell(row=header_row, column=col, value=h))
    widths = [4, 20, 10, 60, 48, 16, 12, 38, 14, 12, 15]
    for col, w in enumerate(widths, start=1):
        ov.column_dimensions[get_column_letter(col)].width = w
    status_dv = DataValidation(type="list", formula1='"Not started,Requested,Received,Uploaded"')
    ov.add_data_validation(status_dv)

    for i, (sheet, kind, key, prio, why, where, target) in enumerate(DATASETS, start=1):
        r = header_row + i
        if kind == "entity":
            _, ncols = entity_sheet(wb, sheet, key)
            filled = f"=COUNTA('{sheet}'!A3:A202)"
            real = "n/a"
        else:
            _, nrows, src_col = param_sheet(wb, sheet, key)
            filled = f"=COUNTA('{sheet}'!A2:A{nrows + 60})"
            real = f"=COUNTIF('{sheet}'!{src_col}2:{src_col}{nrows + 60},\"real\")"
        vals = [i, sheet, prio, why, where, None, None, target, "Not started", filled, real]
        for col, v in enumerate(vals, start=1):
            c = ov.cell(row=r, column=col, value=v)
            c.font = Font(name=FONT, size=10, bold=(col == 2),
                          color={"High": "8C2F1B", "Medium": "7F6000", "Low": "404040"}.get(v, "000000")
                          if col == 3 else "000000")
            c.alignment = Alignment(vertical="top", wrap_text=True)
            c.border = THIN
            if col in (6, 7, 9):
                c.fill = INPUT_FILL
        ov.cell(row=r, column=2).hyperlink = f"#'{sheet}'!A1"
        status_dv.add(f"I{r}")
    last = header_row + len(DATASETS)
    ov.cell(row=last + 2, column=2, value="Datasets uploaded").font = Font(name=FONT, bold=True, size=10)
    ov.cell(row=last + 2, column=3, value=f'=COUNTIF(I{header_row + 1}:I{last},"Uploaded")&" of "&'
                                          f'COUNTA(B{header_row + 1}:B{last})').font = Font(name=FONT, size=10)
    ov.freeze_panes = f"A{header_row + 1}"
    # Excel computes the row counters when the file opens.
    from openpyxl.workbook.properties import CalcProperties
    wb.calculation = CalcProperties(fullCalcOnLoad=True)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


if __name__ == "__main__":
    out = Path(sys.argv[1]) if len(sys.argv) > 1 else Path("templates/FlexiTog_data_request_pack.xlsx")
    print(build(out))
